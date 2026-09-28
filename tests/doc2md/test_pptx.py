# Copyright (c) 2026 PaddlePaddle Authors. All Rights Reserved.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

import pytest

pptx = pytest.importorskip("pptx")
from pptx.util import Inches

from paddleocr._doc2md import convert


def _add_run(paragraph, text, strike=None):
    run = paragraph.add_run()
    run.text = text
    if strike is not None:
        run._r.get_or_add_rPr().set("strike", strike)
    return run


def test_pptx_table_cells(tmp_path):
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table = slide.shapes.add_table(
        2, 2, Inches(1), Inches(1), Inches(4), Inches(2)
    ).table
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Price"
    table.cell(1, 0).text = "Apple"
    para = table.cell(1, 1).text_frame.paragraphs[0]
    _add_run(para, "10", strike="sngStrike")
    _add_run(para, "8")
    path = tmp_path / "table.pptx"
    prs.save(str(path))

    markdown = convert(path).markdown

    assert "<th>Name</th>" in markdown
    assert "<th>Price</th>" in markdown
    assert "<td>Apple</td>" in markdown
    assert "<td><del>10</del>8</td>" in markdown


@pytest.mark.parametrize(
    "strike, expected",
    [("sngStrike", "- ~~old~~"), ("dblStrike", "- ~~old~~"), ("noStrike", "- old")],
)
def test_pptx_text_strikethrough(tmp_path, strike, expected):
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    _add_run(textbox.text_frame.paragraphs[0], "old", strike=strike)
    path = tmp_path / "text.pptx"
    prs.save(str(path))

    markdown = convert(path).markdown

    assert markdown.splitlines() == [expected]
